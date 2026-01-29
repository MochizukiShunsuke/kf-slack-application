import os
import logging
import re
import calendar
from datetime import datetime
from docxtpl import DocxTemplate, InlineImage
from docx import Document
from docx.shared import Emu, Mm
from PIL import Image
import io

from config import (
    ACTIVITY_REPORT_TEMPLATE_ID,
    ACTIVITY_PICTURE_FOLDER_ID,
    MEMBER_PICTURE_FOLDER_ID,
    ACTIVITY_REPORT_CHANNEL_ID,
    ACTIVITY_REPORT_MESSAGE_FOLDER_ID
)

from services.sheet_service import (
    get_activity_report_mentions,
    get_member_department_map
)
from services.drive_service import (
    download_file_to_stream,
    search_files_with_sort,
    search_image_by_name,
    list_files_in_folder,
    read_text_from_docx_stream,
    export_google_doc_text,
    find_folder_by_name
)
from services.slack_service import (
    send_slack_message,
    upload_slack_file
)

logger = logging.getLogger(__name__)

# ---------------------------------------------------------
# Helper Functions
# ---------------------------------------------------------

def get_cell_width_from_doc(doc, tags):
    """テンプレート内の指定タグが存在するセルの幅を取得する"""
    widths = {}
    def scan(parent):
        for table in parent.tables:
            for row in table.rows:
                for cell in row.cells:
                    text = "".join(p.text for p in cell.paragraphs).strip()
                    for tag in tags:
                        if tag in text:
                            widths[tag] = cell.width if cell.width and cell.width > 0 else None
                    scan(cell)
    scan(doc)
    return widths

def get_headline_from_filename(filename):
    """ファイル名からアンダーバー以降を抽出 (例: activity1_大会の様子.jpg -> 大会の様子)"""
    stem = os.path.splitext(filename)[0]
    if "_" in stem:
        return stem.split("_", 1)[1]
    return ""

def get_month_en(month_num):
    """月の英語名を取得"""
    try:
        return calendar.month_name[int(month_num)]
    except:
        return ""

def resize_image_for_docx(image_bytes):
    """画像をWord貼り付け用に変換"""
    try:
        img = Image.open(image_bytes)
        output = io.BytesIO()
        img_format = img.format if img.format else 'JPEG'
        img.save(output, format=img_format)
        output.seek(0)
        return output
    except Exception as e:
        logger.warning(f"Image processing warning: {e}")
        return None

# ---------------------------------------------------------
# UI/Logic: 催促機能
# ---------------------------------------------------------

def run_activity_report_reminder():
    """担当者への催促"""
    try:
        current_month = str(datetime.now().month)
        mentions = get_activity_report_mentions(current_month)
        
        if not mentions:
            msg = f"今月({current_month}月)の担当者が見つかりませんでした。"
            send_slack_message(channel=ACTIVITY_REPORT_CHANNEL_ID, text=msg)
            return msg
        message = (
            f"近活の原稿を作成し、提出してください\n"
            f"{mentions}\n"
            f"提出期限は25日です"
        )
        send_slack_message(channel=ACTIVITY_REPORT_CHANNEL_ID, text=message)
        return f"Successfully sent reminder to: {mentions}"
    except Exception as e:
        logger.error(f"Reminder Error: {e}")
        raise e

# ---------------------------------------------------------
# UI/Logic: 報告書生成機能
# ---------------------------------------------------------

def generate_activity_report(client, channel_id, user_id, target_month):
    temp_file_path = None
    try:
        logger.info(f"Starting report generation for: {target_month}")
        
        # 1. 年月の抽出 (テンプレート埋め込み用)
        year_str = ""
        month_str = ""
        match = re.search(r"(\d{4})年度(\d{1,2})月", target_month)
        if match:
            year_str = match.group(1)
            month_str = match.group(2)
        else:
            # フォールバック
            now = datetime.now() # JSTがあれば JSTで
            month = now.month
            if month >= 10:
                year = now.year + 1
            else:
                year = now.year
            year_str = str(year)
            month_str = str(month)
        
        month_en = get_month_en(month_str)

        # 2. テンプレート準備
        template_stream = download_file_to_stream(ACTIVITY_REPORT_TEMPLATE_ID)
        doc = DocxTemplate(template_stream)
        
        tags = [f'{{{{member{i}}}}}' for i in range(1, 8)] + [f'{{{{activity{i}}}}}' for i in range(1, 7)]
        template_stream.seek(0)
        temp_doc_for_measure = Document(template_stream)
        widths = get_cell_width_from_doc(temp_doc_for_measure, tags)

        # 3. コンテキスト初期化
        context = {
            'year_title': year_str,
            'month_title': month_str,
            'year_side': year_str,
            'month_side': month_en
        }

        # 4. 名簿情報の取得
        member_map = get_member_department_map()

        # =========================================================
        # 5. 親フォルダから対象月のフォルダIDを探す
        # =========================================================
        # 原稿用フォルダ
        message_month_folder_id = find_folder_by_name(ACTIVITY_REPORT_MESSAGE_FOLDER_ID, target_month)
        # 写真用フォルダ
        picture_month_folder_id = find_folder_by_name(ACTIVITY_PICTURE_FOLDER_ID, target_month)

        if not message_month_folder_id or not picture_month_folder_id:
            msg = f"<@{user_id}> ❌ 「{target_month}」のフォルダが見つかりませんでした。\n親フォルダ内に「{target_month}」という名前でフォルダを作成してください。"
            send_slack_message(channel=channel_id, text=msg)
            return

        # =========================================================
        # 6. 原稿ファイルの検索と処理 (対象月フォルダから検索)
        # =========================================================
        search_query = f"'{message_month_folder_id}' in parents and name contains '原稿' and trashed = false"
        input_files = search_files_with_sort(search_query, limit=7)

        for i, file_info in enumerate(input_files, 1):
            fname = file_info['name']
            fid = file_info['id']
            fmime = file_info.get('mimeType', '')
            
            stem = os.path.splitext(fname)[0]
            member_name = stem.split('_')[-1] if "_" in stem else stem
            
            # A. 文章
            text = ""
            if "wordprocessingml.document" in fmime or fname.endswith(".docx"):
                stream = download_file_to_stream(fid)
                text = read_text_from_docx_stream(stream)
            elif "application/vnd.google-apps.document" in fmime:
                text = export_google_doc_text(fid)
            
            # B. 所属
            clean_name = member_name.replace(" ", "").replace("　", "")
            department = member_map.get(clean_name, "所属不明")

            context[f'script{i}'] = text
            context[f'name{i}'] = member_name
            context[f'department{i}'] = department

            # C. 顔写真 (ここは以前のまま、全員分のフォルダから検索)
            photo_stream = search_image_by_name(MEMBER_PICTURE_FOLDER_ID, member_name)
            if photo_stream:
                processed_img = resize_image_for_docx(photo_stream)
                if processed_img:
                    tag_key = f'{{{{member{i}}}}}'
                    w = widths.get(tag_key)
                    img_obj = InlineImage(doc, processed_img, width=Emu(w) if w else Mm(35))
                    context[f'member{i}'] = img_obj
                else:
                    context[f'member{i}'] = ""
            else:
                context[f'member{i}'] = ""

        # =========================================================
        # 7. 活動写真の処理 (対象月フォルダから検索)
        # =========================================================
        files_in_pic_folder = list_files_in_folder(picture_month_folder_id)

        # 【追加】ここデバッグ用ログを入れてください
        # ↓↓↓↓↓↓↓↓↓↓↓↓↓↓↓↓↓↓↓↓↓↓↓↓↓↓↓↓↓↓↓↓↓↓↓
        found_names = [f['name'] for f in files_in_pic_folder]
        logger.info(f"DEBUG: Found files in picture folder: {found_names}")

        for j in range(1, 7):
            prefix = f"activity{j}"
            target_file = next((f for f in files_in_pic_folder if f['name'].startswith(prefix)), None)
            
            headline_key = f'headline{j}'
            img_key = f'activity{j}'
            
            if target_file:
                headline = get_headline_from_filename(target_file['name'])
                context[headline_key] = f"【{headline}】" if headline else ""
                
                p_stream = download_file_to_stream(target_file['id'])
                processed_img = resize_image_for_docx(p_stream)
                if processed_img:
                    tag_key = f'{{{{{img_key}}}}}'
                    w = widths.get(tag_key)
                    img_obj = InlineImage(doc, processed_img, width=Emu(w) if w else Mm(80))
                    context[img_key] = img_obj
                else:
                    context[img_key] = ""
            else:
                context[headline_key] = ""
                context[img_key] = ""

        # 8. レンダリングと保存
        doc.render(context)
        report_title = f"【金沢大学フォーミュラ研究会】{target_month}近況活動報告書"
        filename = f"{report_title}.docx"
        temp_file_path = f"/tmp/{filename}"
        doc.save(temp_file_path)

        # 9. Upload
        with open(temp_file_path, "rb") as f:
            upload_slack_file(
                channel_id=channel_id,
                file_content=f.read(),
                title=report_title,
                filename=filename,
                comment=f"<@{user_id}> {target_month}の近況活動報告書が完成しました！\nご確認をお願いします。 🏎️💨"
            )

    except Exception as e:
        logger.exception("Error in generate_activity_report")
        send_slack_message(
            channel=channel_id,
            text=f"<@{user_id}> ❌ 報告書の作成中にエラーが発生しました: {str(e)}"
        )
    finally:
        if temp_file_path and os.path.exists(temp_file_path):
            try: os.remove(temp_file_path)
            except: pass
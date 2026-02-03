import os
import io
import logging
import google.auth
from google.oauth2.credentials import Credentials
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload
from docx import Document
import requests
from googleapiclient.http import MediaIoBaseDownload, MediaIoBaseUpload


logger = logging.getLogger(__name__)

def get_drive_service():
    try:
        client_id = os.environ.get("GOOGLE_CLIENT_ID")
        client_secret = os.environ.get("GOOGLE_CLIENT_SECRET")
        refresh_token = os.environ.get("GOOGLE_REFRESH_TOKEN")

        if not all([client_id, client_secret, refresh_token]):
            logger.error("❌ OAuth credentials are missing in environment variables.")
            return None

        creds = Credentials(
            None,
            refresh_token=refresh_token,
            token_uri="https://oauth2.googleapis.com/token",
            client_id=client_id,
            client_secret=client_secret
        )

        return build('drive', 'v3', credentials=creds)

    except Exception as e:
        logger.error(f"❌ Failed to create Drive service: {e}")
        return None
    
service = get_drive_service()

def download_file_to_stream(file_id: str) -> io.BytesIO:
    print("drive_service.download_file_to_stream...")
    try:
        service = get_drive_service()
        request = service.files().get_media(fileId=file_id)
        file_stream = io.BytesIO()
        downloader = MediaIoBaseDownload(file_stream, request)
        done = False
        while not done:
            status, done = downloader.next_chunk()
        file_stream.seek(0)
        return file_stream
    except Exception as e:
        logger.exception("Download File To Stream Error")
        raise e

def read_text_from_docx_stream(file_stream: io.BytesIO) -> str:
    print("drive_service.read_text_from_docx_stream...")
    try:
        doc = Document(file_stream)
        full_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text.strip())
        return "\n".join(full_text)
    except Exception:
        logger.exception("Read Text From Docx Stream Error")
        return ""

def list_files_in_folder(folder_id: str):
    print("drive_service.list_files_in_folder...")
    try:
        service = get_drive_service()
        query = f"'{folder_id}' in parents and trashed = false"
        results = service.files().list(
            q=query, 
            fields="files(id, name)",
            pageSize=100
        ).execute()
        return results.get('files', [])
    except Exception:
        logger.exception("List Files In Folder Error")
        return []

def search_files_with_sort(query_str: str, limit=7):
    print("drive_service.search_files_with_sort...")
    try:
        service = get_drive_service()
        results = service.files().list(
            q=query_str,
            fields="files(id, name, mimeType)",
            pageSize=50 
        ).execute()
        
        files = results.get('files', [])
        valid_files = [f for f in files if not f['name'].startswith('~$')]
        sorted_files = sorted(valid_files, key=lambda x: x['name'])
        
        return sorted_files[:limit]
    except Exception:
        logger.exception("Search Files With Sort Error")
        return []

def search_image_by_name(folder_id: str, member_name: str):
    print("drive_service.search_image_by_name...")
    try:
        service = get_drive_service()
        query = f"'{folder_id}' in parents and name contains '{member_name}' and mimeType contains 'image/' and trashed = false"
        results = service.files().list(
            q=query,
            fields="files(id, name)",
            pageSize=1
        ).execute()
        
        files = results.get('files', [])
        if not files:
            return None
            
        return download_file_to_stream(files[0]['id'])
    except Exception:
        logger.exception("Search Image By Name Error")
        return None

def export_google_doc_text(file_id: str) -> str:
    print("drive_service.export_google_doc_text...")
    try:
        service = get_drive_service()
        request = service.files().export_media(fileId=file_id, mimeType='text/plain')
        file_stream = io.BytesIO()
        downloader = MediaIoBaseDownload(file_stream, request)
        done = False
        while not done:
            status, done = downloader.next_chunk()
        return file_stream.getvalue().decode('utf-8')
    except Exception:
        logger.exception("Export Google Doc Text Error")
        return ""
    
def find_folder_by_name(parent_id, folder_name):
    if not service: return None
    try:
        query = f"'{parent_id}' in parents and name = '{folder_name}' and mimeType = 'application/vnd.google-apps.folder' and trashed = false"
        results = service.files().list(q=query, fields="files(id, name)").execute()
        files = results.get('files', [])
        if not files:
            return None
        return files[0]['id']
    except Exception as e:
        logger.error(f"Error finding folder: {e}")
        return None
    
def get_user_draft_text(folder_id: str, user_name: str) -> str:
    print(f"drive_service.get_user_draft_text: {user_name}")
    try:
        service = get_drive_service()
        query = f"'{folder_id}' in parents and name contains '原稿' and name contains '{user_name}' and trashed = false"
        
        results = service.files().list(
            q=query,
            fields="files(id, name, mimeType)",
            pageSize=1
        ).execute()
        
        files = results.get('files', [])
        if not files:
            return ""
            
        file = files[0]
        if "google-apps.document" in file['mimeType']:
            return export_google_doc_text(file['id'])
        else:
            stream = download_file_to_stream(file['id'])
            return read_text_from_docx_stream(stream)
            
    except Exception:
        logger.exception("Get User Draft Text Error")
        return ""

def upload_text_as_docx(folder_id, filename, text_content):
    if not service: return None
    try:
        doc = Document()
        doc.add_paragraph(text_content)
        stream = io.BytesIO()
        doc.save(stream)
        stream.seek(0)
        
        file_metadata = {
            'name': filename,
            'parents': [folder_id]
        }
        media = MediaIoBaseUpload(
            stream,
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            resumable=True
        )

        file = service.files().create(
            body=file_metadata,
            media_body=media,
            fields='id',
            supportsAllDrives=True
        ).execute()
        
        logger.info(f"Uploaded docx: {file.get('id')}")
        return file.get('id')
    except Exception as e:
        logger.error(f"Upload text error: {e}")
        raise e

def upload_image_from_url(folder_id, filename, image_url, token):
    import requests
    if not service: return None
    try:
        headers = {"Authorization": f"Bearer {token}"}
        res = requests.get(image_url, headers=headers)
        
        if res.status_code != 200:
            logger.error(f"Failed to download image from Slack: {res.status_code}")
            return None

        file_metadata = {
            'name': filename,
            'parents': [folder_id]
        }
        
        media = MediaIoBaseUpload(
            io.BytesIO(res.content), 
            mimetype='image/jpeg', 
            resumable=True
        )

        file = service.files().create(
            body=file_metadata,
            media_body=media,
            fields='id',
            supportsAllDrives=True
        ).execute()
        
        logger.info(f"Uploaded image: {file.get('id')}")
        return file.get('id')

    except Exception as e:
        logger.error(f"Upload image error: {e}")
        raise e
    
def upload_docx_from_url(folder_id, filename, docx_url, token):
    if not service: return None
    try:
        headers = {"Authorization": f"Bearer {token}"}
        res = requests.get(docx_url, headers=headers)
        
        if res.status_code != 200:
            logger.error(f"Failed to download docx from Slack: {res.status_code}")
            return None

        file_metadata = {
            'name': filename,
            'parents': [folder_id]
        }
        media = MediaIoBaseUpload(
            io.BytesIO(res.content), 
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document', 
            resumable=True
        )

        file = service.files().create(
            body=file_metadata,
            media_body=media,
            fields='id',
            supportsAllDrives=True
        ).execute()
        
        logger.info(f"Uploaded docx: {file.get('id')}")
        return file.get('id')

    except Exception as e:
        logger.error(f"Upload docx error: {e}")
        raise e